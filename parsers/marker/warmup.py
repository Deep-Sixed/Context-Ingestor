"""Build-time only: fetch every model and font Marker loads lazily.

Marker and Surya download weights (Hugging Face and models.datalab.to) and a
font on first use. Converting a small generated PDF and PPTX here, while the
build still has network, bakes all of it into the image; at run time the
parser has no network and Hugging Face is forced offline.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, "/opt/stele")
import entry  # noqa: E402

HTML = """
<h1>Stele warm-up</h1>
<p>Parsers run in a sandbox with no network.</p>
<table border="1">
  <tr><th>Item</th><th>Count</th></tr>
  <tr><td>Pages</td><td>1</td></tr>
</table>
"""


def main() -> None:
    from pptx import Presentation
    from weasyprint import HTML as Html

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        pdf = root / "warmup.pdf"
        Html(string=HTML).write_pdf(str(pdf))

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Stele warm-up"
        slide.placeholders[1].text = "Sandboxed extraction"
        pptx = root / "warmup.pptx"
        deck.save(str(pptx))

        for doc in (pdf, pptx):
            out = root / f"out-{doc.suffix[1:]}"
            out.mkdir()
            entry.convert(doc, out, dict(entry.DEFAULTS))
            text = (out / "document.md").read_text(encoding="utf-8")
            if "warm-up" not in text:
                raise SystemExit(f"warm-up conversion of {doc.name} produced no text:\n{text}")
            print(f"warm-up: {doc.name} ok")

    entry.MODELS_ROOT.mkdir(parents=True, exist_ok=True)
    (entry.MODELS_ROOT / ".stele-ready").write_text("marker models baked at build time\n")


if __name__ == "__main__":
    main()
