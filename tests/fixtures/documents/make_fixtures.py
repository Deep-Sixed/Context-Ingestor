"""Regenerate the representative documents the parser-image tests parse.

    pip install reportlab pillow python-docx python-pptx
    python tests/fixtures/documents/make_fixtures.py

Each document carries marker phrases (MARKERS below) that a correct
extraction must reproduce; tests/test_parser_images.py checks for them.
The generated files are committed so the tests never need these tools.
"""
from __future__ import annotations

import io
from pathlib import Path

HERE = Path(__file__).resolve().parent

MARKERS = {
    "text_only.pdf": ["Stele fixture: plain text", "evidence before extraction"],
    "scanned.pdf": ["INVOICE", "TOTAL"],
    "tables.pdf": ["Quarterly results", "Revenue", "Widgets", "1250"],
    "multi_column.pdf": ["Left column begins", "Right column begins"],
    "report.docx": ["Stele fixture: word document", "Region", "North", "4200"],
    "slides.pptx": ["Stele fixture: slide deck", "Second slide", "provenance"],
}

LOREM = (
    "Parsers run in a sandbox with no network. Every input is staged and hashed "
    "before extraction, and every output is hashed after it. "
)


def _text_only(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [
        Paragraph("Stele fixture: plain text", styles["Title"]),
        Paragraph("Why evidence before extraction", styles["Heading2"]),
        Paragraph(LOREM * 4, styles["BodyText"]),
        Spacer(1, 12),
        Paragraph(LOREM * 3, styles["BodyText"]),
    ]
    SimpleDocTemplate(str(path), pagesize=letter, title="text only", author="Stele",
                      invariant=1).build(story)


def _scanned(path: Path) -> None:
    """A page that is only an image of text: no text layer, so OCR is needed."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    img = Image.new("L", (1275, 1650), 255)  # letter at 150 dpi
    draw = ImageDraw.Draw(img)
    bold = ImageFont.truetype("DejaVuSans-Bold.ttf", 64)
    body = ImageFont.truetype("DejaVuSans.ttf", 36)
    draw.text((120, 140), "INVOICE", font=bold, fill=0)
    rows = [("Item", "Qty", "Price"), ("Paper", "10", "4.00"), ("Ink", "2", "18.50")]
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            draw.text((120 + j * 360, 320 + i * 70), cell, font=body, fill=0)
    draw.text((120, 620), "TOTAL 77.00", font=bold, fill=0)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.drawImage(ImageReader(buf), 0, 0, width=letter[0], height=letter[1])
    c.showPage()
    c.save()


def _tables(path: Path) -> None:
    """A report page: prose, a captioned table with a plain grid, more prose.

    Kept deliberately page-like: a lone shaded grid on an otherwise empty page
    is classified as a picture by some layout models (Docling), which is not
    what this fixture is meant to test.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    data = [
        ["Product", "Q1", "Q2", "Q3", "Revenue"],
        ["Widgets", "300", "310", "640", "1250"],
        ["Gadgets", "120", "95", "180", "395"],
        ["Gizmos", "80", "82", "90", "252"],
    ]
    table = Table(data, colWidths=[120, 70, 70, 70, 90])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    story = [
        Paragraph("Quarterly results", styles["Title"]),
        Paragraph(LOREM * 2, styles["BodyText"]),
        Paragraph("Units sold per quarter and total revenue are shown in Table 1.",
                  styles["BodyText"]),
        Spacer(1, 12),
        table,
        Spacer(1, 6),
        Paragraph("Table 1: Units sold per quarter and total revenue.", styles["Italic"]),
        Spacer(1, 12),
        Paragraph(LOREM * 3, styles["BodyText"]),
    ]
    SimpleDocTemplate(str(path), pagesize=letter, title="tables", author="Stele",
                      invariant=1).build(story)


def _multi_column(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import BaseDocTemplate, Frame, FrameBreak, PageTemplate, Paragraph

    styles = getSampleStyleSheet()
    width, height = letter
    margin, gap = 54, 18
    col = (width - 2 * margin - gap) / 2
    frames = [
        Frame(margin, margin, col, height - 2 * margin, id="left"),
        Frame(margin + col + gap, margin, col, height - 2 * margin, id="right"),
    ]
    doc = BaseDocTemplate(str(path), pagesize=letter, title="multi column", author="Stele",
                          invariant=1)
    doc.addPageTemplates([PageTemplate(id="two", frames=frames)])
    story = [Paragraph("Left column begins here.", styles["Heading2"])]
    story += [Paragraph(LOREM, styles["BodyText"]) for _ in range(9)]
    story += [FrameBreak(), Paragraph("Right column begins here.", styles["Heading2"])]
    story += [Paragraph(LOREM, styles["BodyText"]) for _ in range(9)]
    doc.build(story)


def _docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Stele"
    d.add_heading("Stele fixture: word document", level=1)
    d.add_paragraph(LOREM)
    table = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Region", "Sales"), ("North", "4200"), ("South", "3100")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    d.add_paragraph("End of document.")
    d.save(str(path))


def _pptx(path: Path) -> None:
    import pptx

    deck = pptx.Presentation()
    deck.core_properties.author = "Stele"
    s1 = deck.slides.add_slide(deck.slide_layouts[0])
    s1.shapes.title.text = "Stele fixture: slide deck"
    s1.placeholders[1].text = "Sandboxed extraction"
    s2 = deck.slides.add_slide(deck.slide_layouts[1])
    s2.shapes.title.text = "Second slide"
    s2.placeholders[1].text = "Every artifact carries its provenance"
    deck.save(str(path))


BUILDERS = {
    "text_only.pdf": _text_only,
    "scanned.pdf": _scanned,
    "tables.pdf": _tables,
    "multi_column.pdf": _multi_column,
    "report.docx": _docx,
    "slides.pptx": _pptx,
}


if __name__ == "__main__":
    for name, build in BUILDERS.items():
        build(HERE / name)
        print(f"wrote {name} ({(HERE / name).stat().st_size} bytes)")
